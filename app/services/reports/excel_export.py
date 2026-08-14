from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def export_excel(report_data, path):
    wb = Workbook()
    wb.remove(wb.active)

    _overview_sheet(wb, report_data["overview"])
    _quality_sheet(wb, report_data["quality"])
    _eda_sheet(wb, report_data["eda"])
    _ml_sheet(wb, report_data["ml_summary"])
    _recommendations_sheet(wb, report_data["recommendations"])

    wb.save(path)


def _style_header(ws, row=1):
    for cell in ws[row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT


def _autosize(ws, max_width=50):
    for col_cells in ws.columns:
        length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=10)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(length + 2, max_width)


def _overview_sheet(wb, overview):
    ws = wb.create_sheet("Overview")
    ws.append(["Metric", "Value"])
    for key, value in overview.items():
        ws.append([key.replace("_", " ").title(), value])
    _style_header(ws)
    _autosize(ws)


def _quality_sheet(wb, quality):
    ws = wb.create_sheet("Data Quality")
    ws.append(["Metric", "Value"])
    ws.append(["Total missing cells", quality["missing"]["total_missing_cells"]])
    ws.append(["Duplicate rows", quality["duplicates"]["count"]])
    ws.append(["Constant columns", ", ".join(quality["constant_columns"]) or "None"])
    ws.append(["High-cardinality columns", ", ".join(quality["high_cardinality_columns"]) or "None"])
    ws.append([])
    ws.append(["Column", "Missing count", "Missing %"])
    for col, info in quality["missing"]["columns_with_missing"].items():
        ws.append([col, info["count"], info["percentage"]])
    ws.append([])
    ws.append(["Column", "Outlier count", "Lower bound", "Upper bound"])
    for col, info in quality["outliers"].items():
        ws.append([col, info["count"], info["lower_bound"], info["upper_bound"]])
    _style_header(ws)
    _autosize(ws)


def _eda_sheet(wb, eda):
    ws = wb.create_sheet("EDA - Numeric")
    ws.append(["Column", "Count", "Mean", "Median", "Std", "Min", "Q1", "Q3", "Max"])
    for col, stats in eda["numeric"].items():
        ws.append([
            col, stats.get("count"), stats.get("mean"), stats.get("median"), stats.get("std"),
            stats.get("min"), stats.get("q1"), stats.get("q3"), stats.get("max"),
        ])
    _style_header(ws)
    _autosize(ws)

    ws2 = wb.create_sheet("EDA - Categorical")
    ws2.append(["Column", "Unique count", "Top value", "Top value count"])
    for col, stats in eda["categorical"].items():
        top = stats.get("top_values", [{}])
        top_val = top[0].get("value") if top else None
        top_count = top[0].get("count") if top else None
        ws2.append([col, stats.get("unique_count"), top_val, top_count])
    _style_header(ws2)
    _autosize(ws2)


def _ml_sheet(wb, ml_summary):
    ws = wb.create_sheet("ML Results")
    best = ml_summary.get("best_model")
    if not best:
        ws.append(["No trained ML model available for this dataset."])
        return

    ws.append(["Best model", best["algorithm"]])
    ws.append(["Task type", best["task_type"]])
    ws.append(["Target column", best["target_column"]])
    ws.append([])
    ws.append(["Metric", "Value"])
    for metric_name, metric_value in best["metrics"].items():
        if metric_name not in ("confusion_matrix", "labels"):
            ws.append([metric_name, metric_value])
    ws.append([])
    ws.append(["Top feature", "Importance %"])
    for f in best.get("top_features", []):
        ws.append([f["feature"], f["importance_pct"]])
    _style_header(ws, row=1)
    _autosize(ws)


def _recommendations_sheet(wb, recommendations):
    ws = wb.create_sheet("Recommendations")
    ws.append(["#", "Recommendation"])
    for i, rec in enumerate(recommendations, start=1):
        ws.append([i, rec])
    _style_header(ws)
    _autosize(ws)
