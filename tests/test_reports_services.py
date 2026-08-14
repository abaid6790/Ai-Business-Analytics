import json
import os

import numpy as np
import pandas as pd
import pytest

from app.services.reports.recommendations import build_recommendations
from app.services.reports.report_data import build_report_data
from app.services.reports.chart_renderer import render_chart_png
from app.services.reports.json_export import export_json
from app.services.reports.csv_export import export_csv
from app.services.reports.excel_export import export_excel
from app.services.reports.pdf_export import export_pdf


class FakeDataset:
    name = "Test Dataset"
    description = "A test dataset"
    original_filename = "test.csv"
    row_count = 100
    column_count = 3
    missing_values_count = 5
    duplicate_rows_count = 2
    file_size_bytes = 4096
    created_at = None
    is_cleaned_copy = False
    project_id = None


class FakeChart:
    def __init__(self, title="Test Chart", spec=None):
        self.title = title
        self.chart_type = "bar"
        self.config_json = json.dumps(spec or {"chart_type": "bar", "x": "region", "y": "amount", "aggregation": "sum"})


class FakeMLModel:
    algorithm = "RandomForestRegressor"
    task_type = "regression"
    target_column = "amount"
    metrics_json = json.dumps({"mae": 5.0, "rmse": 8.0, "r2": 0.9})
    feature_importance_json = json.dumps({"top_features": [{"feature": "units", "importance_pct": 80.0}]})
    explanation_text = "units drives amount."


class FakeAIMessage:
    def __init__(self, content, facts=None):
        self.content = content
        self.computed_context_json = json.dumps(facts) if facts is not None else None


def _sample_df():
    rng = np.random.RandomState(0)
    return pd.DataFrame({
        "region": rng.choice(["north", "south", "east"], 100),
        "amount": rng.normal(500, 100, 100),
        "units": rng.randint(1, 50, 100),
    })


# --- recommendations -----------------------------------------------------
def test_recommendations_flags_duplicates():
    quality = {"duplicates": {"count": 5}, "missing": {"columns_with_missing": {}}, "constant_columns": [], "outliers": {}}
    recs = build_recommendations(quality, {}, [])
    assert any("duplicate" in r.lower() for r in recs)


def test_recommendations_flags_missing_values():
    quality = {
        "duplicates": {"count": 0},
        "missing": {"columns_with_missing": {"age": {"count": 10, "percentage": 25.0}}},
        "constant_columns": [], "outliers": {},
    }
    recs = build_recommendations(quality, {}, [])
    assert any("age" in r and "missing" in r.lower() for r in recs)


def test_recommendations_flags_low_accuracy_model():
    ml_summary = {"best_model": {"algorithm": "LogReg", "task_type": "classification", "metrics": {"accuracy": 0.55}}}
    recs = build_recommendations({}, ml_summary, [])
    assert any("55.0%" in r for r in recs)


def test_recommendations_praises_good_model():
    ml_summary = {"best_model": {"algorithm": "XGBRegressor", "task_type": "regression", "metrics": {"r2": 0.92}}}
    recs = build_recommendations({}, ml_summary, [])
    assert any("ready to use" in r for r in recs)


def test_recommendations_fallback_when_nothing_notable():
    recs = build_recommendations({}, {}, [])
    assert "No significant" in recs[0]


# --- report_data -----------------------------------------------------------
def test_build_report_data_includes_all_sections():
    df = _sample_df()
    data = build_report_data(FakeDataset(), df, [], None, None)
    assert set(data.keys()) == {
        "overview", "quality", "eda", "correlation", "charts",
        "ml_summary", "insights_text", "insight_facts", "recommendations",
    }
    assert data["overview"]["dataset_name"] == "Test Dataset"
    assert data["ml_summary"]["best_model"] is None
    assert data["insights_text"] is None


def test_build_report_data_renders_valid_saved_charts():
    df = _sample_df()
    chart = FakeChart(spec={"chart_type": "bar", "x": "region", "y": "amount", "aggregation": "sum"})
    data = build_report_data(FakeDataset(), df, [chart], None, None)
    assert len(data["charts"]) == 1
    assert data["charts"][0]["title"] == "Test Chart"


def test_build_report_data_skips_broken_saved_chart():
    df = _sample_df()
    broken_chart = FakeChart(spec={"chart_type": "bar", "x": "nonexistent_column"})
    data = build_report_data(FakeDataset(), df, [broken_chart], None, None)
    assert data["charts"] == []  # broken chart silently skipped, not a crash


def test_build_report_data_includes_ml_summary():
    df = _sample_df()
    data = build_report_data(FakeDataset(), df, [], FakeMLModel(), None)
    assert data["ml_summary"]["best_model"]["algorithm"] == "RandomForestRegressor"
    assert data["ml_summary"]["best_model"]["metrics"]["r2"] == 0.9


def test_build_report_data_includes_insights():
    df = _sample_df()
    msg = FakeAIMessage("- Revenue is strong.", facts=[{"type": "trend", "percent_change": 10}])
    data = build_report_data(FakeDataset(), df, [], None, msg)
    assert data["insights_text"] == "- Revenue is strong."
    assert data["insight_facts"][0]["type"] == "trend"


# --- chart_renderer ----------------------------------------------------
def test_render_chart_png_bar():
    png = render_chart_png("Test", {"chart_type": "bar", "labels": ["a", "b"], "datasets": [{"label": "x", "data": [1, 2]}]})
    assert png[:8] == b"\x89PNG\r\n\x1a\n"


def test_render_chart_png_pie():
    png = render_chart_png("Test", {"chart_type": "pie", "labels": ["a", "b"], "datasets": [{"label": "x", "data": [3, 7]}]})
    assert png[:8] == b"\x89PNG\r\n\x1a\n"


def test_render_chart_png_scatter():
    png = render_chart_png("Test", {
        "chart_type": "scatter", "labels": [],
        "datasets": [{"label": "x", "data": [{"x": 1, "y": 2}, {"x": 3, "y": 4}]}],
    })
    assert png[:8] == b"\x89PNG\r\n\x1a\n"


# --- exporters (all 4 formats produce valid, non-empty files) --------------
@pytest.fixture()
def sample_report_data():
    df = _sample_df()
    return build_report_data(FakeDataset(), df, [FakeChart()], FakeMLModel(), FakeAIMessage("- insight"))


def test_export_json_produces_valid_file(sample_report_data, tmp_path):
    path = str(tmp_path / "report.json")
    export_json(sample_report_data, path)
    assert os.path.getsize(path) > 0
    with open(path) as f:
        loaded = json.load(f)
    assert loaded["overview"]["dataset_name"] == "Test Dataset"


def test_export_csv_produces_valid_file(sample_report_data, tmp_path):
    path = str(tmp_path / "report.csv")
    export_csv(sample_report_data, path)
    assert os.path.getsize(path) > 0
    with open(path) as f:
        content = f.read()
    assert "Overview" in content
    assert "Recommendations" in content


def test_export_excel_produces_valid_workbook(sample_report_data, tmp_path):
    path = str(tmp_path / "report.xlsx")
    export_excel(sample_report_data, path)
    assert os.path.getsize(path) > 0

    from openpyxl import load_workbook
    wb = load_workbook(path)
    assert "Overview" in wb.sheetnames
    assert "Data Quality" in wb.sheetnames
    assert "ML Results" in wb.sheetnames


def test_export_pdf_produces_valid_multipage_pdf(sample_report_data, tmp_path):
    path = str(tmp_path / "report.pdf")
    export_pdf(sample_report_data, path, "Test Report Title")
    assert os.path.getsize(path) > 0

    from pypdf import PdfReader
    reader = PdfReader(path)
    assert len(reader.pages) >= 3
    first_page_text = reader.pages[0].extract_text()
    assert "Test Report Title" in first_page_text


def test_export_pdf_handles_no_charts_no_ml_no_insights(tmp_path):
    df = _sample_df()
    data = build_report_data(FakeDataset(), df, [], None, None)
    path = str(tmp_path / "minimal.pdf")
    export_pdf(data, path, "Minimal Report")
    assert os.path.getsize(path) > 0
